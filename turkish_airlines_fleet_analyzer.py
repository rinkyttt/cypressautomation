#!/usr/bin/env python3
"""
Turkish Airlines Fleet Analyzer
Fetches flight data for ALL Turkish Airlines aircraft using FlightRadar24 API
Provides detailed statistics per aircraft registration including total flight time and flight count
"""

import requests
import json
import csv
import time
from datetime import datetime, timedelta
from typing import List, Dict, Optional, Tuple
import os
from dataclasses import dataclass, asdict
from collections import defaultdict
import pandas as pd
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils.dataframe import dataframe_to_rows


@dataclass
class FlightData:
    """Data structure for flight information"""
    registration: str
    flight_number: str
    departure_airport: str
    arrival_airport: str
    departure_time: str
    arrival_time: str
    flight_duration: str
    aircraft_type: str
    status: str
    date: str
    duration_minutes: int = 0


@dataclass
class AircraftStats:
    """Statistics for a single aircraft"""
    registration: str
    aircraft_type: str
    total_flights: int
    total_flight_time_minutes: int
    total_flight_time_hours: float
    average_flight_time_minutes: float
    first_flight_date: str
    last_flight_date: str
    unique_routes: int
    most_common_route: str
    most_common_route_count: int


class FlightRadar24API:
    """FlightRadar24 API client for fetching aircraft and flight data"""
    
    def __init__(self, api_key: str):
        self.api_key = api_key
        self.base_url = "https://api.flightradar24.com"
        self.headers = {
            "Authorization": f"Bearer {api_key}",
            "Content-Type": "application/json",
            "User-Agent": "Turkish-Airlines-Fleet-Analyzer/1.0"
        }
        
    def get_airline_fleet(self, airline_code: str = "TK") -> List[Dict]:
        """
        Retrieve fleet information for Turkish Airlines (ALL aircraft types)
        
        Args:
            airline_code: IATA code for Turkish Airlines (TK)
            
        Returns:
            List of all aircraft in the fleet
        """
        url = f"{self.base_url}/common/v1/airline/fleet.json"
        params = {"airline": airline_code}
        
        try:
            response = requests.get(url, headers=self.headers, params=params)
            response.raise_for_status()
            data = response.json()
            
            fleet = data.get("aircraft", [])
            
            # Group by aircraft type for summary
            type_counts = defaultdict(int)
            for aircraft in fleet:
                aircraft_type = aircraft.get("type", "Unknown")
                type_counts[aircraft_type] += 1
            
            print(f"Found {len(fleet)} aircraft in Turkish Airlines fleet")
            print("Aircraft types in fleet:")
            for aircraft_type, count in sorted(type_counts.items()):
                print(f"  {aircraft_type}: {count} aircraft")
            
            return fleet
            
        except requests.exceptions.RequestException as e:
            print(f"Error fetching fleet data: {e}")
            return []
    
    def get_aircraft_flights(self, registration: str, aircraft_type: str, days_back: int = 30) -> List[FlightData]:
        """
        Retrieve flight history for a specific aircraft
        
        Args:
            registration: Aircraft registration number (e.g., TC-LTI)
            aircraft_type: Aircraft type code (e.g., A21N, B77W)
            days_back: Number of days to look back for flight history
            
        Returns:
            List of FlightData objects
        """
        url = f"{self.base_url}/common/v1/flight/list.json"
        params = {
            "registration": registration,
            "from": int((datetime.now() - timedelta(days=days_back)).timestamp()),
            "to": int(datetime.now().timestamp())
        }
        
        flights = []
        
        try:
            response = requests.get(url, headers=self.headers, params=params)
            response.raise_for_status()
            data = response.json()
            
            for flight in data.get("result", []):
                departure_ts = flight.get("time_scheduled_departure")
                arrival_ts = flight.get("time_scheduled_arrival")
                duration_minutes = self._calculate_duration_minutes(departure_ts, arrival_ts)
                
                flight_data = FlightData(
                    registration=registration,
                    flight_number=flight.get("flight", ""),
                    departure_airport=flight.get("airport_origin", {}).get("iata", ""),
                    arrival_airport=flight.get("airport_destination", {}).get("iata", ""),
                    departure_time=self._format_timestamp(departure_ts),
                    arrival_time=self._format_timestamp(arrival_ts),
                    flight_duration=self._format_duration(duration_minutes),
                    aircraft_type=aircraft_type,
                    status=flight.get("status", ""),
                    date=self._format_date(departure_ts),
                    duration_minutes=duration_minutes
                )
                flights.append(flight_data)
                
            print(f"Retrieved {len(flights)} flights for {aircraft_type} {registration}")
            return flights
            
        except requests.exceptions.RequestException as e:
            print(f"Error fetching flights for {registration}: {e}")
            return []
    
    def _format_timestamp(self, timestamp: Optional[int]) -> str:
        """Convert Unix timestamp to readable format"""
        if not timestamp:
            return ""
        return datetime.fromtimestamp(timestamp).strftime("%Y-%m-%d %H:%M:%S")
    
    def _format_date(self, timestamp: Optional[int]) -> str:
        """Convert Unix timestamp to date format"""
        if not timestamp:
            return ""
        return datetime.fromtimestamp(timestamp).strftime("%Y-%m-%d")
    
    def _calculate_duration_minutes(self, departure: Optional[int], arrival: Optional[int]) -> int:
        """Calculate flight duration in minutes from departure and arrival timestamps"""
        if not departure or not arrival:
            return 0
        
        duration_seconds = arrival - departure
        return max(0, duration_seconds // 60)  # Convert to minutes
    
    def _format_duration(self, minutes: int) -> str:
        """Format duration from minutes to HH:MM format"""
        if minutes <= 0:
            return ""
        
        hours = minutes // 60
        mins = minutes % 60
        return f"{hours:02d}:{mins:02d}"


class TurkishAirlinesFleetAnalyzer:
    """Main class for analyzing Turkish Airlines entire fleet"""
    
    def __init__(self, api_key: str):
        self.api = FlightRadar24API(api_key)
        self.all_flights = []
        self.aircraft_stats = {}
        
    def analyze_fleet(self, days_back: int = 30) -> Dict[str, AircraftStats]:
        """
        Analyze entire Turkish Airlines fleet
        
        Args:
            days_back: Number of days to look back for flight history
            
        Returns:
            Dictionary of aircraft statistics by registration
        """
        print(f"Analyzing Turkish Airlines fleet for the past {days_back} days...")
        fleet = self.api.get_airline_fleet("TK")
        
        if not fleet:
            print("No aircraft found in fleet or API error occurred")
            return {}
        
        all_flights = []
        aircraft_stats = {}
        
        for i, aircraft in enumerate(fleet, 1):
            registration = aircraft.get("registration", "")
            aircraft_type = aircraft.get("type", "Unknown")
            
            print(f"Processing aircraft {i}/{len(fleet)}: {aircraft_type} {registration}")
            
            flights = self.api.get_aircraft_flights(registration, aircraft_type, days_back)
            all_flights.extend(flights)
            
            # Calculate statistics for this aircraft
            if flights:
                stats = self._calculate_aircraft_stats(registration, aircraft_type, flights)
                aircraft_stats[registration] = stats
            else:
                # Create empty stats for aircraft with no flights
                aircraft_stats[registration] = AircraftStats(
                    registration=registration,
                    aircraft_type=aircraft_type,
                    total_flights=0,
                    total_flight_time_minutes=0,
                    total_flight_time_hours=0.0,
                    average_flight_time_minutes=0.0,
                    first_flight_date="",
                    last_flight_date="",
                    unique_routes=0,
                    most_common_route="",
                    most_common_route_count=0
                )
            
            # Rate limiting - wait between requests
            time.sleep(1)
        
        self.all_flights = all_flights
        self.aircraft_stats = aircraft_stats
        
        print(f"\nAnalysis complete!")
        print(f"Total aircraft analyzed: {len(aircraft_stats)}")
        print(f"Total flights retrieved: {len(all_flights)}")
        
        return aircraft_stats
    
    def _calculate_aircraft_stats(self, registration: str, aircraft_type: str, flights: List[FlightData]) -> AircraftStats:
        """Calculate comprehensive statistics for a single aircraft"""
        if not flights:
            return AircraftStats(
                registration=registration,
                aircraft_type=aircraft_type,
                total_flights=0,
                total_flight_time_minutes=0,
                total_flight_time_hours=0.0,
                average_flight_time_minutes=0.0,
                first_flight_date="",
                last_flight_date="",
                unique_routes=0,
                most_common_route="",
                most_common_route_count=0
            )
        
        # Calculate flight time statistics
        total_minutes = sum(flight.duration_minutes for flight in flights)
        total_hours = total_minutes / 60.0
        avg_minutes = total_minutes / len(flights) if flights else 0
        
        # Find date range
        dates = [flight.date for flight in flights if flight.date]
        first_date = min(dates) if dates else ""
        last_date = max(dates) if dates else ""
        
        # Analyze routes
        routes = defaultdict(int)
        for flight in flights:
            if flight.departure_airport and flight.arrival_airport:
                route = f"{flight.departure_airport}-{flight.arrival_airport}"
                routes[route] += 1
        
        most_common_route = ""
        most_common_count = 0
        if routes:
            most_common_route, most_common_count = max(routes.items(), key=lambda x: x[1])
        
        return AircraftStats(
            registration=registration,
            aircraft_type=aircraft_type,
            total_flights=len(flights),
            total_flight_time_minutes=total_minutes,
            total_flight_time_hours=total_hours,
            average_flight_time_minutes=avg_minutes,
            first_flight_date=first_date,
            last_flight_date=last_date,
            unique_routes=len(routes),
            most_common_route=most_common_route,
            most_common_route_count=most_common_count
        )
    
    def export_aircraft_stats_to_csv(self, filename: str = "turkish_airlines_aircraft_stats.csv"):
        """Export aircraft statistics to CSV file"""
        if not self.aircraft_stats:
            print("No aircraft statistics to export")
            return
        
        with open(filename, 'w', newline='', encoding='utf-8') as csvfile:
            fieldnames = [
                'registration', 'aircraft_type', 'total_flights', 
                'total_flight_time_hours', 'total_flight_time_minutes',
                'average_flight_time_minutes', 'first_flight_date', 
                'last_flight_date', 'unique_routes', 'most_common_route',
                'most_common_route_count'
            ]
            
            writer = csv.DictWriter(csvfile, fieldnames=fieldnames)
            writer.writeheader()
            
            for stats in self.aircraft_stats.values():
                writer.writerow(asdict(stats))
        
        print(f"Aircraft statistics exported to {filename}")
    
    def export_flights_to_csv(self, filename: str = "turkish_airlines_all_flights.csv"):
        """Export all flight data to CSV file"""
        if not self.all_flights:
            print("No flight data to export")
            return
        
        with open(filename, 'w', newline='', encoding='utf-8') as csvfile:
            fieldnames = [
                'registration', 'aircraft_type', 'flight_number', 
                'departure_airport', 'arrival_airport', 'departure_time', 
                'arrival_time', 'flight_duration', 'duration_minutes',
                'status', 'date'
            ]
            
            writer = csv.DictWriter(csvfile, fieldnames=fieldnames)
            writer.writeheader()
            
            for flight in self.all_flights:
                writer.writerow(asdict(flight))
        
        print(f"Flight data exported to {filename}")
    
    def export_to_json(self, stats_filename: str = "aircraft_stats.json", flights_filename: str = "all_flights.json"):
        """Export data to JSON files"""
        if self.aircraft_stats:
            stats_dicts = {reg: asdict(stats) for reg, stats in self.aircraft_stats.items()}
            with open(stats_filename, 'w', encoding='utf-8') as jsonfile:
                json.dump(stats_dicts, jsonfile, indent=2, ensure_ascii=False)
            print(f"Aircraft statistics exported to {stats_filename}")
        
        if self.all_flights:
            flight_dicts = [asdict(flight) for flight in self.all_flights]
            with open(flights_filename, 'w', encoding='utf-8') as jsonfile:
                json.dump(flight_dicts, jsonfile, indent=2, ensure_ascii=False)
            print(f"Flight data exported to {flights_filename}")
    
    def export_to_excel(self, filename: str = "turkish_airlines_fleet_analysis.xlsx"):
        """Export comprehensive data to Excel with multiple worksheets and formatting"""
        if not self.aircraft_stats and not self.all_flights:
            print("No data to export")
            return
        
        # Create workbook
        wb = Workbook()
        
        # Remove default sheet
        wb.remove(wb.active)
        
        # Define styles
        header_font = Font(bold=True, color="FFFFFF")
        header_fill = PatternFill(start_color="366092", end_color="366092", fill_type="solid")
        border = Border(
            left=Side(style='thin'),
            right=Side(style='thin'),
            top=Side(style='thin'),
            bottom=Side(style='thin')
        )
        center_align = Alignment(horizontal="center", vertical="center")
        
        # 1. Aircraft Statistics Sheet
        if self.aircraft_stats:
            ws_stats = wb.create_sheet("Aircraft Statistics")
            
            # Convert to DataFrame for easier manipulation
            stats_data = []
            for stats in self.aircraft_stats.values():
                stats_data.append({
                    'Registration': stats.registration,
                    'Aircraft Type': stats.aircraft_type,
                    'Total Flights': stats.total_flights,
                    'Total Flight Hours': round(stats.total_flight_time_hours, 1),
                    'Total Flight Minutes': stats.total_flight_time_minutes,
                    'Avg Flight Time (min)': round(stats.average_flight_time_minutes, 1),
                    'First Flight Date': stats.first_flight_date,
                    'Last Flight Date': stats.last_flight_date,
                    'Unique Routes': stats.unique_routes,
                    'Most Common Route': stats.most_common_route,
                    'Top Route Count': stats.most_common_route_count
                })
            
            df_stats = pd.DataFrame(stats_data)
            
            # Sort by total flight hours descending
            df_stats = df_stats.sort_values('Total Flight Hours', ascending=False)
            
            # Add headers
            for r in dataframe_to_rows(df_stats, index=False, header=True):
                ws_stats.append(r)
            
            # Format headers
            for cell in ws_stats[1]:
                cell.font = header_font
                cell.fill = header_fill
                cell.alignment = center_align
                cell.border = border
            
            # Format data cells
            for row in ws_stats.iter_rows(min_row=2):
                for cell in row:
                    cell.border = border
                    if cell.column in [3, 4, 5, 6, 9, 11]:  # Numeric columns
                        cell.alignment = Alignment(horizontal="center")
            
            # Auto-adjust column widths
            for column in ws_stats.columns:
                max_length = 0
                column_letter = column[0].column_letter
                for cell in column:
                    try:
                        if len(str(cell.value)) > max_length:
                            max_length = len(str(cell.value))
                    except:
                        pass
                adjusted_width = min(max_length + 2, 20)
                ws_stats.column_dimensions[column_letter].width = adjusted_width
        
        # 2. Aircraft Type Summary Sheet
        if self.aircraft_stats:
            ws_types = wb.create_sheet("Aircraft Type Summary")
            
            # Calculate type statistics
            type_stats = defaultdict(lambda: {'count': 0, 'flights': 0, 'hours': 0, 'active': 0})
            for stats in self.aircraft_stats.values():
                aircraft_type = stats.aircraft_type
                type_stats[aircraft_type]['count'] += 1
                type_stats[aircraft_type]['flights'] += stats.total_flights
                type_stats[aircraft_type]['hours'] += stats.total_flight_time_hours
                if stats.total_flights > 0:
                    type_stats[aircraft_type]['active'] += 1
            
            # Create type summary data
            type_data = []
            for aircraft_type, data in sorted(type_stats.items()):
                avg_flights = data['flights'] / data['active'] if data['active'] > 0 else 0
                avg_hours = data['hours'] / data['active'] if data['active'] > 0 else 0
                type_data.append({
                    'Aircraft Type': aircraft_type,
                    'Total Aircraft': data['count'],
                    'Active Aircraft': data['active'],
                    'Total Flights': data['flights'],
                    'Total Hours': round(data['hours'], 1),
                    'Avg Flights/Aircraft': round(avg_flights, 1),
                    'Avg Hours/Aircraft': round(avg_hours, 1)
                })
            
            df_types = pd.DataFrame(type_data)
            df_types = df_types.sort_values('Total Hours', ascending=False)
            
            # Add to worksheet
            for r in dataframe_to_rows(df_types, index=False, header=True):
                ws_types.append(r)
            
            # Format headers
            for cell in ws_types[1]:
                cell.font = header_font
                cell.fill = header_fill
                cell.alignment = center_align
                cell.border = border
            
            # Format data cells
            for row in ws_types.iter_rows(min_row=2):
                for cell in row:
                    cell.border = border
                    if cell.column > 1:  # Numeric columns
                        cell.alignment = Alignment(horizontal="center")
            
            # Auto-adjust column widths
            for column in ws_types.columns:
                max_length = 0
                column_letter = column[0].column_letter
                for cell in column:
                    try:
                        if len(str(cell.value)) > max_length:
                            max_length = len(str(cell.value))
                    except:
                        pass
                adjusted_width = min(max_length + 2, 20)
                ws_types.column_dimensions[column_letter].width = adjusted_width
        
        # 3. Top Performers Sheet
        if self.aircraft_stats:
            ws_top = wb.create_sheet("Top Performers")
            
            # Top 20 by flights
            top_flights = sorted(self.aircraft_stats.values(), key=lambda x: x.total_flights, reverse=True)[:20]
            top_hours = sorted(self.aircraft_stats.values(), key=lambda x: x.total_flight_time_hours, reverse=True)[:20]
            
            # Add section headers and data
            ws_top.append(["TOP 20 AIRCRAFT BY FLIGHT COUNT"])
            ws_top.append(["Rank", "Registration", "Aircraft Type", "Total Flights", "Total Hours", "Avg Hours/Flight"])
            
            for i, stats in enumerate(top_flights, 1):
                avg_flight_hours = stats.total_flight_time_hours / stats.total_flights if stats.total_flights > 0 else 0
                ws_top.append([
                    i, stats.registration, stats.aircraft_type, 
                    stats.total_flights, round(stats.total_flight_time_hours, 1), 
                    round(avg_flight_hours, 1)
                ])
            
            # Add spacing
            ws_top.append([])
            ws_top.append([])
            
            # Top by hours
            ws_top.append(["TOP 20 AIRCRAFT BY FLIGHT HOURS"])
            ws_top.append(["Rank", "Registration", "Aircraft Type", "Total Hours", "Total Flights", "Unique Routes"])
            
            for i, stats in enumerate(top_hours, 1):
                ws_top.append([
                    i, stats.registration, stats.aircraft_type,
                    round(stats.total_flight_time_hours, 1), stats.total_flights, stats.unique_routes
                ])
            
            # Format the sheet
            # Section headers
            ws_top['A1'].font = Font(bold=True, size=14)
            ws_top['A24'].font = Font(bold=True, size=14)
            
            # Table headers
            for cell in ws_top[2]:
                if cell.value:
                    cell.font = header_font
                    cell.fill = header_fill
                    cell.alignment = center_align
                    cell.border = border
            
            for cell in ws_top[25]:
                if cell.value:
                    cell.font = header_font
                    cell.fill = header_fill
                    cell.alignment = center_align
                    cell.border = border
            
            # Auto-adjust column widths
            for column in ws_top.columns:
                max_length = 0
                column_letter = column[0].column_letter
                for cell in column:
                    try:
                        if len(str(cell.value)) > max_length:
                            max_length = len(str(cell.value))
                    except:
                        pass
                adjusted_width = min(max_length + 2, 15)
                ws_top.column_dimensions[column_letter].width = adjusted_width
        
        # 4. All Flights Sheet (if data exists)
        if self.all_flights:
            ws_flights = wb.create_sheet("All Flights")
            
            # Convert flights to DataFrame
            flights_data = []
            for flight in self.all_flights:
                flights_data.append({
                    'Registration': flight.registration,
                    'Aircraft Type': flight.aircraft_type,
                    'Flight Number': flight.flight_number,
                    'Departure': flight.departure_airport,
                    'Arrival': flight.arrival_airport,
                    'Departure Time': flight.departure_time,
                    'Arrival Time': flight.arrival_time,
                    'Duration': flight.flight_duration,
                    'Duration (min)': flight.duration_minutes,
                    'Status': flight.status,
                    'Date': flight.date
                })
            
            df_flights = pd.DataFrame(flights_data)
            
            # Sort by date and departure time
            df_flights = df_flights.sort_values(['Date', 'Departure Time'])
            
            # Add to worksheet
            for r in dataframe_to_rows(df_flights, index=False, header=True):
                ws_flights.append(r)
            
            # Format headers
            for cell in ws_flights[1]:
                cell.font = header_font
                cell.fill = header_fill
                cell.alignment = center_align
                cell.border = border
            
            # Auto-adjust column widths
            for column in ws_flights.columns:
                max_length = 0
                column_letter = column[0].column_letter
                for cell in column:
                    try:
                        if len(str(cell.value)) > max_length:
                            max_length = len(str(cell.value))
                    except:
                        pass
                adjusted_width = min(max_length + 2, 25)
                ws_flights.column_dimensions[column_letter].width = adjusted_width
        
        # Save the workbook
        wb.save(filename)
        print(f"Excel file exported to {filename}")
        
        # Print summary of what was exported
        sheet_names = [sheet.title for sheet in wb.worksheets]
        print(f"Excel file contains {len(sheet_names)} worksheets: {', '.join(sheet_names)}")
    
    def print_detailed_summary(self):
        """Print comprehensive fleet analysis summary"""
        if not self.aircraft_stats:
            print("No aircraft statistics available")
            return
        
        print("\n" + "="*80)
        print("TURKISH AIRLINES FLEET ANALYSIS - DETAILED SUMMARY")
        print("="*80)
        
        # Overall fleet statistics
        total_aircraft = len(self.aircraft_stats)
        active_aircraft = len([s for s in self.aircraft_stats.values() if s.total_flights > 0])
        total_flights = sum(stats.total_flights for stats in self.aircraft_stats.values())
        total_flight_hours = sum(stats.total_flight_time_hours for stats in self.aircraft_stats.values())
        
        print(f"Total Aircraft in Fleet: {total_aircraft}")
        print(f"Active Aircraft (with flights): {active_aircraft}")
        print(f"Total Flights: {total_flights:,}")
        print(f"Total Flight Hours: {total_flight_hours:,.1f}")
        print(f"Average Flight Hours per Aircraft: {total_flight_hours/active_aircraft:.1f}" if active_aircraft > 0 else "N/A")
        
        # Aircraft type breakdown
        print(f"\n{'AIRCRAFT TYPE SUMMARY':^80}")
        print("-" * 80)
        type_stats = defaultdict(lambda: {'count': 0, 'flights': 0, 'hours': 0})
        
        for stats in self.aircraft_stats.values():
            aircraft_type = stats.aircraft_type
            type_stats[aircraft_type]['count'] += 1
            type_stats[aircraft_type]['flights'] += stats.total_flights
            type_stats[aircraft_type]['hours'] += stats.total_flight_time_hours
        
        print(f"{'Type':<8} {'Count':<6} {'Flights':<8} {'Hours':<10} {'Avg Hours/Aircraft':<15}")
        print("-" * 80)
        for aircraft_type, data in sorted(type_stats.items()):
            avg_hours = data['hours'] / data['count'] if data['count'] > 0 else 0
            print(f"{aircraft_type:<8} {data['count']:<6} {data['flights']:<8} {data['hours']:<10.1f} {avg_hours:<15.1f}")
        
        # Top performers
        print(f"\n{'TOP PERFORMING AIRCRAFT':^80}")
        print("-" * 80)
        
        # Most flights
        most_flights = sorted(self.aircraft_stats.values(), key=lambda x: x.total_flights, reverse=True)[:10]
        print(f"\nTop 10 by Flight Count:")
        print(f"{'Registration':<12} {'Type':<8} {'Flights':<8} {'Hours':<10} {'Avg/Flight':<10}")
        print("-" * 60)
        for stats in most_flights:
            avg_flight = stats.average_flight_time_minutes / 60 if stats.average_flight_time_minutes > 0 else 0
            print(f"{stats.registration:<12} {stats.aircraft_type:<8} {stats.total_flights:<8} "
                  f"{stats.total_flight_time_hours:<10.1f} {avg_flight:<10.1f}")
        
        # Most flight hours
        most_hours = sorted(self.aircraft_stats.values(), key=lambda x: x.total_flight_time_hours, reverse=True)[:10]
        print(f"\nTop 10 by Flight Hours:")
        print(f"{'Registration':<12} {'Type':<8} {'Hours':<10} {'Flights':<8} {'Routes':<8}")
        print("-" * 60)
        for stats in most_hours:
            print(f"{stats.registration:<12} {stats.aircraft_type:<8} {stats.total_flight_time_hours:<10.1f} "
                  f"{stats.total_flights:<8} {stats.unique_routes:<8}")
        
        # Detailed per-aircraft breakdown
        print(f"\n{'DETAILED AIRCRAFT BREAKDOWN':^80}")
        print("-" * 80)
        print(f"{'Registration':<12} {'Type':<8} {'Flights':<8} {'Hours':<10} {'Routes':<8} {'Top Route':<15}")
        print("-" * 80)
        
        for registration in sorted(self.aircraft_stats.keys()):
            stats = self.aircraft_stats[registration]
            top_route = stats.most_common_route[:14] if stats.most_common_route else "None"
            print(f"{stats.registration:<12} {stats.aircraft_type:<8} {stats.total_flights:<8} "
                  f"{stats.total_flight_time_hours:<10.1f} {stats.unique_routes:<8} {top_route:<15}")


def main():
    """Main function to run the fleet analyzer"""
    # Get API key from environment variable or prompt user
    api_key = os.getenv('FLIGHTRADAR24_API_KEY')
    
    if not api_key:
        print("Please set your FlightRadar24 API key:")
        print("1. As environment variable: export FLIGHTRADAR24_API_KEY='your_api_key'")
        print("2. Or enter it when prompted below:")
        api_key = input("Enter your FlightRadar24 API key: ").strip()
    
    if not api_key:
        print("API key is required to use this script")
        return
    
    # Initialize analyzer
    analyzer = TurkishAirlinesFleetAnalyzer(api_key)
    
    # Get user preferences
            try:
            days_back = int(input("Enter number of days to look back (default 30): ") or "30")
        except ValueError:
            days_back = 30
    
    print(f"\nAnalyzing Turkish Airlines fleet for the last {days_back} days...")
    print("This may take a while as we're processing the entire fleet...")
    
    # Analyze fleet
    aircraft_stats = analyzer.analyze_fleet(days_back)
    
    if aircraft_stats:
        # Print detailed summary
        analyzer.print_detailed_summary()
        
        # Export data
        print(f"\n{'EXPORTING DATA':^80}")
        print("-" * 80)
        analyzer.export_aircraft_stats_to_csv()
        analyzer.export_flights_to_csv()
        analyzer.export_to_json()
        analyzer.export_to_excel()
        
        print("\nFleet analysis complete!")
        print("Files created:")
        print("- turkish_airlines_aircraft_stats.csv (per-aircraft statistics)")
        print("- turkish_airlines_all_flights.csv (all flight records)")
        print("- aircraft_stats.json (statistics in JSON format)")
        print("- all_flights.json (flight records in JSON format)")
        print("- turkish_airlines_fleet_analysis.xlsx (comprehensive Excel workbook)")
    else:
        print("No fleet data was retrieved. Please check your API key and try again.")


if __name__ == "__main__":
    main()